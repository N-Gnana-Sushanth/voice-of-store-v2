"""
Build SignalForge_Voice_of_Store_Round1.pptx
Official Round 1 Submission Deck (7 Slides)
Author: SignalForge (Team Leader: Nadagavani Gnana Sushanth)
"""

import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# --- COLOR PALETTE (Voice of Store Visual Identity) ---
COLOR_BG = RGBColor(255, 255, 255)            # Pure White
COLOR_NAVY = RGBColor(15, 32, 66)             # Deep Navy (#0F2042)
COLOR_TEXT_DARK = RGBColor(17, 24, 39)        # Slate Dark Text (#111827)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)    # Slate Muted Text (#64748B)
COLOR_BLUE = RGBColor(37, 99, 235)            # Royal Blue Accent (#2563EB)
COLOR_BLUE_LIGHT = RGBColor(239, 246, 255)    # Very Light Blue Tint (#EFF6FF)
COLOR_CYAN = RGBColor(2, 132, 199)            # Cyan Accent (#0284C7)
COLOR_TEAL = RGBColor(13, 148, 136)           # Deep Teal (#0D9488)
COLOR_TEAL_LIGHT = RGBColor(240, 253, 250)    # Light Mint/Teal Tint (#F0FDFA)
COLOR_PURPLE = RGBColor(124, 58, 237)         # Purple Accent (#7C3AED)
COLOR_PURPLE_LIGHT = RGBColor(245, 243, 255)  # Light Purple Tint (#F5F3FF)
COLOR_CARD_BG = RGBColor(248, 250, 252)       # Light Slate Card BG (#F8FAFC)
COLOR_BORDER = RGBColor(226, 232, 240)        # Subtle Border (#E2E8F0)
COLOR_BORDER_BLUE = RGBColor(191, 219, 254)   # Subtle Blue Border (#BFDBFE)
COLOR_AMBER = RGBColor(217, 119, 6)           # Warning Amber (#D97706)
COLOR_AMBER_LIGHT = RGBColor(254, 243, 199)   # Light Amber (#FEF3C7)
COLOR_RED = RGBColor(220, 38, 38)             # Danger Red (#DC2626)
COLOR_RED_LIGHT = RGBColor(254, 242, 242)     # Light Red (#FEF2F2)
COLOR_GREEN = RGBColor(22, 163, 74)           # Success Green (#16A34A)
COLOR_GREEN_LIGHT = RGBColor(240, 253, 244)   # Light Green (#F0FDF4)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"

def set_shape_flat(shape, fill_color, border_color=None, border_width_pt=1):
    """Sets flat solid fill and subtle outline on a shape."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width_pt)
    else:
        shape.line.fill.background()

def add_header(slide, title_text, category_text, subtitle_text=None):
    """Creates a standardized modern header block across slides 2-7."""
    # Eyebrow / Category
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.7), Inches(0.25))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.name = FONT_HEADING
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_BLUE

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.7), Inches(0.48))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_NAVY

    # Subtitle
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.18), Inches(11.7), Inches(0.3))
        tf_s = sub_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle_text
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(11)
        p_s.font.color.rgb = COLOR_TEXT_MUTED

def add_card(slide, left, top, width, height, fill_color=COLOR_CARD_BG, border_color=COLOR_BORDER, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    """Adds a rounded card container."""
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    set_shape_flat(shape, fill_color, border_color, 1)
    return shape

def add_badge(slide, left, top, width, height, text, bg_color, text_color, font_size=9):
    """Adds a compact status badge or tag."""
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    set_shape_flat(badge, bg_color, None)
    tf = badge.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_HEADING
    p.font.size = Pt(font_size)
    p.font.bold = True
    p.font.color.rgb = text_color
    return badge

print("Helper definitions loaded.")
