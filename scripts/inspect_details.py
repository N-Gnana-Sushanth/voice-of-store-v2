import os
import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def inspect_all():
    prs = pptx.Presentation(r'C:\Users\user\Downloads\VoiceOfStore_Round1.pptx')
    print("=== PRESENTATION PROPERTIES ===")
    print(f"Slide width: {prs.slide_width / 914400:.2f} inches ({prs.slide_width} EMU)")
    print(f"Slide height: {prs.slide_height / 914400:.2f} inches ({prs.slide_height} EMU)")
    print(f"Slide count: {len(prs.slides)}")
    
    print("\n=== SLIDE 1 FULL SHAPES DUMP ===")
    slide1 = prs.slides[0]
    for i, s in enumerate(slide1.shapes):
        text_content = ""
        if s.has_text_frame:
            text_content = " | ".join([p.text.strip() for p in s.text_frame.paragraphs if p.text.strip()])
        fill_info = "no fill"
        try:
            if s.fill.type:
                fill_info = f"fill_type={s.fill.type}"
                if hasattr(s.fill, 'fore_color') and s.fill.fore_color.type:
                    fill_info += f", rgb={s.fill.fore_color.rgb}"
        except:
            pass
        print(f"[{i:02d}] {s.name:<18} type={s.shape_type:<18} L={s.left/914400:.2f}\" T={s.top/914400:.2f}\" W={s.width/914400:.2f}\" H={s.height/914400:.2f}\" {fill_info} Text: '{text_content}'")

if __name__ == '__main__':
    inspect_all()
