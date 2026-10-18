import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')
prs = pptx.Presentation(r'C:\Users\user\Downloads\VoiceOfStore_Round1.pptx')

for i in range(1, len(prs.slides)):
    slide = prs.slides[i]
    print(f"\n=================== SLIDE {i+1} ===================")
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    print(f"  {txt}")
        elif s.has_table:
            print("  [TABLE]")
            for r in s.table.rows:
                print("    " + " | ".join([c.text.strip().replace('\n', ' ') for c in r.cells]))
