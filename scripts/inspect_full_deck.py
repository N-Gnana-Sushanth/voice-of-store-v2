import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation(r'C:\Users\user\Downloads\VoiceOfStore_Round1.pptx')

print(f"Presentation has {len(prs.slides)} slides.")
print(f"Dimensions: {prs.slide_width} x {prs.slide_height} (16:9 widescreen = 12192000 x 6858000)")

slide1 = prs.slides[0]
print("\n--- SLIDE 1 SHAPE DETAILS ---")
for i, s in enumerate(slide1.shapes):
    txt = ""
    font_details = ""
    if s.has_text_frame:
        txt = " // ".join([p.text.strip() for p in s.text_frame.paragraphs if p.text.strip()])
        if s.text_frame.paragraphs:
            p0 = s.text_frame.paragraphs[0]
            if p0.runs:
                r0 = p0.runs[0]
                color = str(r0.font.color.rgb) if (r0.font and r0.font.color and hasattr(r0.font.color, 'rgb') and r0.font.color.rgb) else "no-rgb"
                font_details = f"font={r0.font.name}, sz={r0.font.size.pt if r0.font.size else 'none'}, bold={r0.font.bold}, color={color}"
            elif p0.font:
                color = str(p0.font.color.rgb) if (p0.font.color and hasattr(p0.font.color, 'rgb') and p0.font.color.rgb) else "no-rgb"
                font_details = f"font={p0.font.name}, sz={p0.font.size.pt if p0.font.size else 'none'}, bold={p0.font.bold}, color={color}"
    
    fill_col = "no-fill"
    try:
        if s.fill.type:
            fill_col = f"type={s.fill.type}"
            if hasattr(s.fill, 'fore_color') and s.fill.fore_color and hasattr(s.fill.fore_color, 'rgb'):
                fill_col += f", rgb={s.fill.fore_color.rgb}"
    except:
        pass
    
    print(f"Shape {i:02d}: name='{s.name}' L={s.left/914400:.2f} T={s.top/914400:.2f} W={s.width/914400:.2f} H={s.height/914400:.2f} | {fill_col} | {font_details} | Text: '{txt}'")
