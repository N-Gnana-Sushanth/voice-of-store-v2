import pptx
import sys

def inspect_slide1():
    prs = pptx.Presentation(r'C:\Users\user\Downloads\VoiceOfStore_Round1.pptx')
    slide1 = prs.slides[0]
    print(f"Slide 1 shapes count: {len(slide1.shapes)}")
    for i, s in enumerate(slide1.shapes):
        print(f"Shape {i}: name='{s.name}', type={s.shape_type}, left={s.left}, top={s.top}, width={s.width}, height={s.height}")
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    font_name = p.font.name if p.font else None
                    font_size = p.font.size.pt if p.font and p.font.size else None
                    print(f"    P: '{p_text}' (font={font_name}, size={font_size})")

if __name__ == '__main__':
    inspect_slide1()
